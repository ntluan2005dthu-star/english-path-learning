# ============================================================
# modules/excel_store.py — Đọc/ghi Excel an toàn
# ============================================================
import os
import shutil
import tempfile
from datetime import date, datetime
from threading import Lock

from openpyxl import Workbook, load_workbook

DATA_DIR = 'data'
DATA_FILE = os.path.join(DATA_DIR, 'english_learning_data.xlsx')
BACKUP_DIR = os.path.join(DATA_DIR, 'backups')

_LOCK = Lock()


SCHEMAS = {
    'Vocabulary': ['id', 'word', 'ipa', 'pos', 'vi_meaning', 'en_definition',
                   'collocation', 'example_en', 'example_vi', 'topic', 'cefr',
                   'difficulty', 'source_name', 'source_url', 'added_at',
                   'status', 'favorite', 'next_review', 'interval_days', 'ease'],
    'Phrases': ['id', 'phrase_en', 'phrase_vi', 'topic', 'source_name', 'source_url'],
    'Grammar': ['id', 'title_vi', 'title_en', 'content_vi', 'example_en',
                'example_vi', 'level', 'source_name', 'source_url'],
    'Lessons': ['id', 'month', 'title_vi', 'title_en', 'skill', 'description_vi',
                'content_vi', 'estimated_minutes', 'unlock_order', 'vocab_topics'],
    'Exercises': ['id', 'lesson_id', 'type', 'question', 'options', 'answer',
                  'explanation_vi', 'knowledge_point', 'vocab_id'],
    'TestResults': ['id', 'date', 'quiz_type', 'lesson_id', 'total', 'correct',
                    'score', 'duration_sec', 'details'],
    'StudyProgress': ['id', 'date', 'tasks_done', 'new_words', 'reviewed_words',
                      'quiz_score', 'minutes', 'note'],
    'Mistakes': ['id', 'qid', 'question', 'question_type', 'user_answer',
                 'correct_answer', 'explanation_vi', 'knowledge_point',
                 'wrong_count', 'last_practiced', 'resolved'],
    'Sources': ['id', 'name', 'url', 'type', 'verified'],
    'Settings': ['id', 'key', 'value', 'note'],
}


def _row_from_dict(sheet_name, row_dict, row_id):
    headers = SCHEMAS[sheet_name]
    values = []
    for h in headers:
        if h == 'id':
            values.append(row_id)
        else:
            v = row_dict.get(h, '')
            if isinstance(v, datetime):
                v = v.strftime('%Y-%m-%d')
            elif isinstance(v, date):
                v = v.strftime('%Y-%m-%d')
            elif v is None:
                v = ''
            values.append(v)
    return values


def _ensure_dirs():
    os.makedirs(DATA_DIR, exist_ok=True)
    os.makedirs(BACKUP_DIR, exist_ok=True)


def ensure_workbook():
    """Tạo file Excel nếu chưa có, kèm dữ liệu mẫu."""
    _ensure_dirs()
    if os.path.exists(DATA_FILE):
        return
    from modules import seed_data

    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, headers in SCHEMAS.items():
        ws = wb.create_sheet(sheet_name)
        ws.append(headers)

    # Seed Vocabulary
    ws = wb['Vocabulary']
    for i, (word, ipa, pos, vi, en_def, colloc, ex_en, ex_vi, topic,
            cefr, diff, src, url) in enumerate(seed_data.VOCABULARY, start=1):
        ws.append([i, word, ipa or 'Chưa xác minh', pos, vi, en_def, colloc,
                   ex_en, ex_vi, topic, cefr, diff, src, url,
                   seed_data.TODAY, 'new', 0, '', 0, 2.5])

    # Seed Phrases
    ws = wb['Phrases']
    for i, (en, vi, topic, src, url) in enumerate(seed_data.PHRASES, start=1):
        ws.append([i, en, vi, topic, src, url])

    # Seed Grammar
    ws = wb['Grammar']
    for i, (t_vi, t_en, content, ex_en, ex_vi, level, src, url) in enumerate(seed_data.GRAMMAR, start=1):
        ws.append([i, t_vi, t_en, content, ex_en, ex_vi, level, src, url])

    # Seed Lessons
    ws = wb['Lessons']
    for i, (month, t_vi, t_en, skill, desc, content, mins, order, topics) in enumerate(seed_data.LESSONS, start=1):
        ws.append([i, month, t_vi, t_en, skill, desc, content, mins, order, topics])

    # Seed Exercises
    ws = wb['Exercises']
    for i, (lesson_id, qtype, question, options, answer, expl, kp, vid) in enumerate(seed_data.EXERCISES, start=1):
        ws.append([i, lesson_id, qtype, question, options, answer, expl, kp, vid])

    # Seed Sources
    ws = wb['Sources']
    for i, (name, url, typ, verified) in enumerate(seed_data.SOURCES, start=1):
        ws.append([i, name, url, typ, verified])

    # Seed Settings
    ws = wb['Settings']
    for i, (key, value) in enumerate(seed_data.SETTINGS, start=1):
        ws.append([i, key, value, ''])

    _atomic_save(wb)


def _backup_current():
    if os.path.exists(DATA_FILE):
        _ensure_dirs()
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        backup_path = os.path.join(BACKUP_DIR, f'backup_{ts}.xlsx')
        try:
            shutil.copy2(DATA_FILE, backup_path)
        except OSError:
            pass


def _atomic_save(wb):
    """Lưu workbook qua file tạm rồi thay thế (atomic)."""
    _ensure_dirs()
    fd, tmp_path = tempfile.mkstemp(suffix='.xlsx', dir=DATA_DIR)
    os.close(fd)
    try:
        wb.save(tmp_path)
        os.replace(tmp_path, DATA_FILE)
    except Exception:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass
        raise


def _load():
    if not os.path.exists(DATA_FILE):
        ensure_workbook()
    return load_workbook(DATA_FILE)


def _sheet_to_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = list(rows[0])
    out = []
    for row in rows[1:]:
        if all(c is None or str(c).strip() == '' for c in row):
            continue
        d = {}
        for h, v in zip(headers, row):
            if h is None:
                continue
            if isinstance(v, (datetime, date)):
                v = v.strftime('%Y-%m-%d')
            elif v is None:
                v = ''
            d[h] = v
        out.append(d)
    return out


def read_sheet(name):
    if name not in SCHEMAS:
        raise ValueError(f'Sheet không tồn tại: {name}')
    with _LOCK:
        wb = _load()
        if name not in wb.sheetnames:
            return []
        rows = _sheet_to_rows(wb[name])
    return rows


def _next_id(ws):
    max_id = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        try:
            val = int(row[0])
            if val > max_id:
                max_id = val
        except (TypeError, ValueError, IndexError):
            continue
    return max_id + 1


def append_row(name, row_dict):
    if name not in SCHEMAS:
        raise ValueError(f'Sheet không tồn tại: {name}')
    with _LOCK:
        wb = _load()
        ws = wb[name]
        new_id = _next_id(ws)
        values = _row_from_dict(name, row_dict, new_id)
        ws.append(values)
        _backup_current()
        _atomic_save(wb)
    row_dict = dict(row_dict)
    row_dict['id'] = new_id
    return row_dict


def update_row(name, row_id, changes):
    if name not in SCHEMAS:
        raise ValueError(f'Sheet không tồn tại: {name}')
    with _LOCK:
        wb = _load()
        ws = wb[name]
        headers = [c.value for c in ws[1]]
        target_row = None
        for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                if int(row[0]) == int(row_id):
                    target_row = r_idx
                    break
            except (TypeError, ValueError, IndexError):
                continue
        if target_row is None:
            raise ValueError(f'Không tìm thấy id={row_id} trong {name}')
        for col_idx, h in enumerate(headers, start=1):
            if h in changes:
                v = changes[h]
                if isinstance(v, (datetime, date)):
                    v = v.strftime('%Y-%m-%d')
                ws.cell(row=target_row, column=col_idx, value=v)
        _backup_current()
        _atomic_save(wb)


def delete_row(name, row_id):
    if name not in SCHEMAS:
        raise ValueError(f'Sheet không tồn tại: {name}')
    with _LOCK:
        wb = _load()
        ws = wb[name]
        for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            try:
                if int(row[0]) == int(row_id):
                    ws.delete_rows(r_idx, 1)
                    _backup_current()
                    _atomic_save(wb)
                    return True
            except (TypeError, ValueError, IndexError):
                continue
    return False